"""
Script to create a PowerPoint presentation for EDWH to EDP Migration Success Story
Informatica EMEA Master Class Session
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def create_title_slide(prs):
    """Create the title slide"""
    slide_layout = prs.slide_layouts[0]  # Title slide layout
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "EDWH to EDP Migration Success Story"
    subtitle.text = "Informatica Power Center to IDMC\nTeradata to BigQuery\n\nInformatica EMEA Master Class"
    
    return slide

def create_content_slide(prs, title, content_points):
    """Create a content slide with bullet points"""
    slide_layout = prs.slide_layouts[1]  # Title and Content layout
    slide = prs.slides.add_slide(slide_layout)
    
    title_shape = slide.shapes.title
    title_shape.text = title
    
    body_shape = slide.placeholders[1]
    text_frame = body_shape.text_frame
    text_frame.clear()
    
    for point in content_points:
        if isinstance(point, tuple):
            # Main point with sub-points
            p = text_frame.add_paragraph()
            p.text = point[0]
            p.level = 0
            for sub_point in point[1]:
                p = text_frame.add_paragraph()
                p.text = sub_point
                p.level = 1
        else:
            # Simple point
            p = text_frame.add_paragraph()
            p.text = point
            p.level = 0
    
    return slide

def main():
    # Create presentation
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Title Slide
    create_title_slide(prs)
    
    # Slide 2: Introduction about Sunrise
    create_content_slide(prs, "Introduction - Sunrise", [
        "Company Overview",
        "Business Domain and Operations",
        "Digital Transformation Journey",
        "Technology Stack Evolution",
        "Current Market Position"
    ])
    
    # Slide 3: Introduction about EDWH
    create_content_slide(prs, "Introduction - EDWH (Enterprise Data Warehouse)", [
        "Legacy Architecture Overview",
        ("Technology Stack:", [
            "Informatica PowerCenter for ETL",
            "Teradata as Data Warehouse Platform",
            "On-premise Infrastructure"
        ]),
        "Data Processing Capabilities",
        "Business Intelligence Integration",
        "Architecture Landscape Components"
    ])
    
    # Slide 4: EDWH Architecture Landscape
    create_content_slide(prs, "EDWH Architecture Landscape", [
        "Source Systems Integration",
        "ETL Processing Layer (Informatica PowerCenter)",
        "Data Storage Layer (Teradata)",
        "Data Consumption Layer",
        "Monitoring and Operations",
        "Security and Governance Framework"
    ])
    
    # Slide 5: Migration Needs
    create_content_slide(prs, "Migration Needs - Why Migration Was Needed", [
        "Cost Optimization Requirements",
        "Scalability Limitations of Legacy System",
        "Need for Cloud-Native Capabilities",
        "Improved Performance and Processing Speed",
        "Modern Data Architecture Requirements",
        "Business Agility and Time-to-Market",
        "Maintenance and Support Challenges",
        "Innovation and Future-Readiness"
    ])
    
    # Slide 6: Challenges Faced
    create_content_slide(prs, "Challenges Faced During Migration", [
        "Complex Data Dependencies",
        "Business Continuity Requirements",
        "Data Quality and Validation",
        "Skill Gap in New Technologies",
        "Change Management and User Adoption",
        "Migration Timeline Constraints",
        "Integration with Existing Systems",
        "Risk Management and Rollback Planning"
    ])
    
    # Slide 7: Solution - Part 1 (Assessment & Planning)
    create_content_slide(prs, "Solution - Assessment & Planning", [
        ("Assessment Phase:", [
            "Complete inventory of mappings and workflows",
            "Complexity analysis",
            "Dependency mapping",
            "Resource planning"
        ]),
        ("Cloudification Strategy:", [
            "Cloud architecture design",
            "Security and compliance framework",
            "Migration approach selection",
            "Cost modeling"
        ])
    ])
    
    # Slide 8: Solution - Part 2 (Development & Testing)
    create_content_slide(prs, "Solution - Development & Testing", [
        ("Parallel Development:", [
            "Dual environment setup",
            "Agile development methodology",
            "Code migration and refactoring",
            "Testing in isolated environments"
        ]),
        ("DVT (Data Validation Tool):", [
            "Automated data comparison",
            "Quality assurance",
            "Reconciliation reporting",
            "Issue tracking and resolution"
        ])
    ])
    
    # Slide 9: Solution - Part 3 (Tools & Technologies)
    create_content_slide(prs, "Solution - Tools & Technologies", [
        ("Google API for Query Conversion:", [
            "Teradata to BigQuery SQL translation",
            "Automated query optimization",
            "Syntax compatibility handling"
        ]),
        ("Environments & Parallel Run:", [
            "Development, Testing, UAT, Production",
            "Parallel execution monitoring",
            "Performance benchmarking",
            "Cutover planning and execution"
        ]),
        "IDMC (Informatica Intelligent Data Management Cloud)",
        "BigQuery as Modern Data Warehouse"
    ])
    
    # Slide 10: Benefits Realized
    create_content_slide(prs, "Benefits Realized", [
        ("Processing Efficiency:", [
            "Faster data processing times",
            "Improved query performance",
            "Enhanced scalability",
            "Real-time capabilities"
        ]),
        ("Cost Savings:", [
            "Reduced infrastructure costs",
            "Pay-as-you-go pricing model",
            "Lower maintenance overhead",
            "Optimized resource utilization"
        ]),
        ("PDO (Push Down Optimization) on All Mappings:", [
            "Leveraging BigQuery compute power",
            "Reduced data movement",
            "Improved performance"
        ])
    ])
    
    # Slide 11: Future Outlook - Part 1
    create_content_slide(prs, "Future Outlook - Innovations", [
        ("CLAIRGPT Integration:", [
            "AI-powered data insights",
            "Automated documentation",
            "Intelligent data discovery",
            "Natural language querying"
        ]),
        ("Additional Utilities & Tools:", [
            "Advanced monitoring solutions",
            "Self-service analytics",
            "Data cataloging enhancements",
            "Automated data quality checks"
        ])
    ])
    
    # Slide 12: Future Outlook - Part 2 (IDMC Topics)
    create_content_slide(prs, "Future Outlook - Topics to Address with IDMC", [
        ("Operational Topics:", [
            "Metadata management optimization",
            "Salesforce merge transformations",
            "Enhanced monitoring and alerting",
            "Performance tuning"
        ]),
        ("Process Improvements:", [
            "Better assessment methodologies",
            "Parallel execution of pipelines",
            "Improved deployment processes"
        ]),
        ("DevOps Integration:", [
            "Bitbucket integration",
            "CI/CD pipeline automation",
            "Version control best practices",
            "Automated deployment workflows"
        ])
    ])
    
    # Slide 13: Conclusion
    create_content_slide(prs, "Conclusion", [
        "Successful migration from EDWH to EDP completed",
        "Significant improvements in performance and cost",
        "Modern cloud-native architecture established",
        "Strong foundation for future innovations",
        "Continuous improvement mindset",
        "Ready for next-generation analytics",
        "",
        "Thank You!",
        "Questions & Discussion"
    ])
    
    # Save the presentation
    output_file = 'EDWH_to_EDP_Migration_Success_Story.pptx'
    prs.save(output_file)
    print(f"Presentation created successfully: {output_file}")
    
    # Print summary
    print(f"\nPresentation Summary:")
    print(f"Total Slides: {len(prs.slides)}")
    print(f"File: {output_file}")

if __name__ == "__main__":
    main()
